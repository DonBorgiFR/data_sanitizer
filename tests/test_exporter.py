# -*- coding: utf-8 -*-
"""
Unit tests for the SAP Excel exporter, ZIP packager, and deduplication blocking performance.
"""
import io
import zipfile
import time
import pandas as pd
import openpyxl
from sanitizer.exporter import generate_sap_staging_excel
from sanitizer.ui_actions import create_migration_zip
from sanitizer.dedupe import find_duplicates
from sanitizer.sap_templates import SAP_TEMPLATES

def test_sap_excel_generation():
    """
    Validates that the generated Excel contains the 3-row headers and reorders/filters
    the columns according to the SAP template, omitting technical columns.
    """
    # Sample data containing official fields and technical/log columns
    data = {
        "PARTNER_ID": ["BP001", "BP002"],
        "NAME": ["Empresa A", "Empresa B"],
        "TAX_ID": ["ESA12345678", "ESB87654321"],
        "POSTAL_CODE": ["08770", "28001"],
        "PROVINCE": ["BARCELONA", "MADRID"],
        "COUNTRY": ["ES", "ES"],
        "STREET": ["Calle Industria 10", "Paseo de la Castellana 20"],
        "NAME_VALID": [True, True],
        "TAX_ID_ERROR": ["", ""],
        "_CLEAN_LOG": ["Clean", "Clean"],
        "UNMAPPED_FIELD": ["XYZ", "123"]
    }
    df = pd.DataFrame(data)
    
    # Generate the Excel bytes for Business Partner template
    excel_bytes = generate_sap_staging_excel(df, "Business Partner (BP)")
    
    # Load workbook from memory
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    
    # The official Business Partner template has exactly 10 columns
    official_cols = list(SAP_TEMPLATES["Business Partner (BP)"]["columns"].keys())
    assert ws.max_column == 10
    
    # Check Row 1 (Technical Names)
    for col_idx, col_name in enumerate(official_cols, start=1):
        assert ws.cell(row=1, column=col_idx).value == col_name
        
    # Check Row 2 (Descriptions)
    assert ws.cell(row=2, column=2).value == "Nombre de la empresa o persona física" # NAME
    
    # Check Row 3 (Format Metada)
    # PARTNER_ID is required, length 10 -> C(10) *
    assert ws.cell(row=3, column=1).value == "C(10) *"
    # IBAN is optional, length 34 -> C(34)
    assert ws.cell(row=3, column=4).value == "C(34)"
    
    # Check data rows start at row 4
    assert ws.cell(row=4, column=1).value == "BP001"
    assert ws.cell(row=4, column=2).value == "Empresa A"
    assert ws.cell(row=5, column=1).value == "BP002"
    assert ws.cell(row=5, column=2).value == "Empresa B"
    
    # Verify that the technical columns and unmapped fields are NOT in the excel
    for col_idx in range(1, ws.max_column + 1):
        tech_header = ws.cell(row=1, column=col_idx).value
        assert tech_header not in ("NAME_VALID", "TAX_ID_ERROR", "_CLEAN_LOG", "UNMAPPED_FIELD")

def test_migration_zip_packaging():
    """
    Validates that the ZIP archive packages all files (Excel load, Excel audit, HTML, TXT logs, and CSVs).
    Nota de evolucion: La expectativa cambio por evolucion deliberada del producto entre Sesion 8 y Sesion 11
    (los archivos Excel/control de Sesion 8 ahora son condicionales y el log se renombro a auditoria_ejecucion.log).
    """
    data = {
        "PARTNER_ID": ["BP001"],
        "NAME": ["Empresa A"],
        "TAX_ID": ["ESA12345678"]
    }
    df = pd.DataFrame(data)
    html_report = "<html><body>Readiness Report Mock</body></html>"
    logs = [
        {"timestamp": "12:00:00", "level": "SYSTEM", "message": "Proceso iniciado"},
        {"timestamp": "12:00:05", "level": "DATA", "message": "Datos cargados"}
    ]
    
    # Case A: With Costing and Inventory Results (Optional legacy annexes should be included)
    costing_results = {
        "cost_df": pd.DataFrame({"MATERIAL": ["MAT01"], "TOTAL_COST": [10.5]})
    }
    inventory_results = {
        "df": pd.DataFrame({"MATERIAL": ["MAT01"], "STOCK": [100]})
    }
    
    zip_bytes = create_migration_zip(
        main_df=df,
        template_name="Business Partner (BP)",
        costing_results=costing_results,
        inventory_results=inventory_results,
        html_report=html_report,
        logs=logs
    )
    
    # Verify files inside ZIP for Case A
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as z:
        files = z.namelist()
        # Session 11 Mandatory files
        assert "datos_limpios.csv" in files
        assert "registros_rechazados.csv" in files
        assert "readiness_report.html" in files
        assert "auditoria_ejecucion.log" in files
        
        # Session 8 Optional annexes (present because costing/inventory have data)
        assert "staging_sap_load.xlsx" in files
        assert "full_audit_sanitized.xlsx" in files
        assert "costing_simulation.csv" in files
        assert "inventory_health.csv" in files
        
        # Verify CSV files encoding (UTF-8-sig has BOM \xef\xbb\xbf) and separator ';'
        costing_csv_bytes = z.read("costing_simulation.csv")
        assert costing_csv_bytes.startswith(b'\xef\xbb\xbf')
        costing_csv_str = costing_csv_bytes.decode("utf-8-sig")
        assert ";" in costing_csv_str.split("\n")[0]
        
        inventory_csv_bytes = z.read("inventory_health.csv")
        assert inventory_csv_bytes.startswith(b'\xef\xbb\xbf')
        inventory_csv_str = inventory_csv_bytes.decode("utf-8-sig")
        assert ";" in inventory_csv_str.split("\n")[0]
        
        # Verify log contents
        log_content = z.read("auditoria_ejecucion.log").decode("utf-8")
        assert "DATASANITIZER ERP MIGRATION ENGINE" in log_content
        
        # Verify html contents
        html_content = z.read("readiness_report.html").decode("utf-8")
        assert "Readiness Report Mock" in html_content
        
    # Case B: Without Costing and Inventory Results (Optional legacy annexes should NOT be included)
    zip_bytes_simple = create_migration_zip(
        main_df=df,
        template_name="Business Partner (BP)",
        costing_results=None,
        inventory_results=None,
        html_report=html_report,
        logs=logs
    )
    
    with zipfile.ZipFile(io.BytesIO(zip_bytes_simple)) as z:
        files = z.namelist()
        # Session 11 Mandatory files must be present
        assert "datos_limpios.csv" in files
        assert "registros_rechazados.csv" in files
        assert "readiness_report.html" in files
        assert "auditoria_ejecucion.log" in files
        
        # Session 8 Optional annexes must NOT be present
        assert "staging_sap_load.xlsx" not in files
        assert "full_audit_sanitized.xlsx" not in files
        assert "costing_simulation.csv" not in files
        assert "inventory_health.csv" not in files

def test_deduplication_blocking_performance():
    """
    Validates deduplication accuracy under blocking keys, demonstrating that 
    the blocking method does not lose obvious duplicates and runs successfully.
    """
    # Create synthetic dataset with repeating patterns
    names = []
    tax_ids = []
    post_codes = []
    
    # Base records
    base_names = ["INDUSTRIAS METALICAS GOMEZ SL", "CONSTRUCCIONES SANZ SA", "ALIMENTACION MARTINEZ S.L."]
    base_taxes = ["ESA12345678", "ESB87654321", "ESC11223344"]
    base_cps = ["08770", "28001", "41002"]
    
    # Generate 400 clean records, inserting duplicates
    for i in range(400):
        base_idx = i % 3
        # Small variations to create potential duplicates
        if i < 15:
            # First 15 are close matches
            names.append(base_names[base_idx] + f" V{i}")
        else:
            names.append(f"EMPRESA INDEPENDIENTE N{i}")
            
        tax_ids.append(base_taxes[base_idx] if i < 10 else f"ESA{10000000 + i}")
        post_codes.append(base_cps[base_idx])
        
    df = pd.DataFrame({
        "NAME": names,
        "TAX_ID": tax_ids,
        "POSTAL_CODE": post_codes
    })
    
    # Run exhaustive (none)
    t0 = time.time()
    dupes_none = find_duplicates(
        df, 
        text_column="NAME", 
        tax_id_column="TAX_ID", 
        postal_code_column="POSTAL_CODE", 
        threshold=85.0, 
        blocking_method="none"
    )
    time_none = time.time() - t0
    
    # Run with blocking (first_3_chars)
    t1 = time.time()
    dupes_blocked = find_duplicates(
        df, 
        text_column="NAME", 
        tax_id_column="TAX_ID", 
        postal_code_column="POSTAL_CODE", 
        threshold=85.0, 
        blocking_method="first_3_chars"
    )
    time_blocked = time.time() - t1
    
    # 1. Assert that obvious duplicates are not lost (compare matches found for key records)
    matches_none = sum(len(g["matches"]) for g in dupes_none)
    matches_blocked = sum(len(g["matches"]) for g in dupes_blocked)
    assert matches_none > 0
    assert matches_blocked == matches_none
    
    # 2. Assert no exception and correct execution
    assert isinstance(dupes_blocked, list)
    
    # Log information (will show up in stdout if running manually or on failure)
    print(f"Exhaustive time: {time_none:.4f}s | Blocked time: {time_blocked:.4f}s")
